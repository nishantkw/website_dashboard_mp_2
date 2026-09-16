"""Build combined NHA + Dashboard server requirement Excel (costs blank for manual fill)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def style_range(ws, start_row, end_row, start_col, end_col, **kwargs):
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(r, c)
            for k, v in kwargs.items():
                setattr(cell, k, v)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Cost Estimate"

    thin = Border(
        left=Side(style="thin", color="666666"),
        right=Side(style="thin", color="666666"),
        top=Side(style="thin", color="666666"),
        bottom=Side(style="thin", color="666666"),
    )
    header_fill = PatternFill("solid", fgColor="003366")
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    title_font = Font(name="Calibri", bold=True, size=14, color="003366")
    subtitle_font = Font(name="Calibri", bold=True, size=11)
    normal = Font(name="Calibri", size=11)
    total_fill = PatternFill("solid", fgColor="FFFF00")
    total_font = Font(name="Calibri", bold=True, size=11)
    section_fill = PatternFill("solid", fgColor="D9E8F5")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Column widths (same layout as original)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16

    # ---------- Section 1: NHA AWS Server ----------
    ws.merge_cells("A1:E1")
    ws["A1"] = "NHA AWS Server"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = section_fill

    ws.merge_cells("A2:E2")
    ws["A2"] = "Cost Estimate of Server for Data Replication (costs to be filled by team)"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["#", "Item", "Specification", "Monthly Cost", "Annual Cost"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(3, col, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    nha_rows = [
        (1, "Virtual Machine", "8vCPU and 32GB RAM, Linux OS, Postgre DB"),
        (2, "OS Disk", "256 GB"),
        (3, "Data Disk", "2048 GB"),
        (4, "IP Address", "1 Static IP"),
        (5, "Egress", "200 GB"),
        (6, "Defender", ""),
    ]

    r = 4
    for num, item, spec in nha_rows:
        ws.cell(r, 1, num).font = normal
        ws.cell(r, 2, item).font = normal
        ws.cell(r, 3, spec).font = normal
        ws.cell(r, 4, "").font = normal  # blank — team fills
        ws.cell(r, 5, "").font = normal
        for c in range(1, 6):
            ws.cell(r, c).border = thin
            ws.cell(r, c).alignment = center if c in (1, 4, 5) else left
        r += 1

    # Total row
    ws.cell(r, 1, 7).font = total_font
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(r, 2, "Total Cost").font = total_font
    ws.cell(r, 4, "").font = total_font
    ws.cell(r, 5, "").font = total_font
    for c in range(1, 6):
        ws.cell(r, c).fill = total_fill
        ws.cell(r, c).border = thin
        ws.cell(r, c).alignment = center if c != 2 else left

    # ---------- Section 2: Dashboard Server Requirement ----------
    r += 3
    dash_title_row = r
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(r, 1, "Dashboard Server Requirement").font = title_font
    ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(r, 1).fill = section_fill

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(
        r,
        1,
        "PMJAY / Website Dashboard — App hosting (~50 users; DB on client existing server; costs to be filled by team)",
    ).font = subtitle_font
    ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")

    r += 1
    for col, h in enumerate(headers, 1):
        cell = ws.cell(r, col, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    dash_rows = [
        (
            1,
            "Backend API server",
            "1 server, 8 vCPU / 16 GB RAM (Node.js / Express) — EC2 or ECS Fargate",
        ),
        (
            2,
            "Frontend hosting",
            "Static React build + CDN (Amazon S3 + CloudFront)",
        ),
        (
            3,
            "SSL certificate",
            "Free public HTTPS certificate (AWS Certificate Manager)",
        ),
        (
            4,
            "Domain + DNS",
            "1 domain (.in / .com) + DNS hosting (Route 53 / registrar)",
        ),
        (
            5,
            "Cache",
            "~1.5 GB Redis (ElastiCache cache.t4g.small) for KPI caching",
        ),
        (
            6,
            "File / export storage",
            "Amazon S3 — uploads, Excel exports, file backups",
        ),
        (
            7,
            "Monitoring & logs",
            "Amazon CloudWatch — application logs + basic alarms",
        ),
        (
            8,
            "Internet data transfer",
            "AWS data transfer / CDN egress — light usage (~50 users)",
        ),
    ]

    r += 1
    first_dash_data = r
    for num, item, spec in dash_rows:
        ws.cell(r, 1, num).font = normal
        ws.cell(r, 2, item).font = normal
        ws.cell(r, 3, spec).font = normal
        ws.cell(r, 4, "").font = normal
        ws.cell(r, 5, "").font = normal
        for c in range(1, 6):
            ws.cell(r, c).border = thin
            ws.cell(r, c).alignment = center if c in (1, 4, 5) else left
        r += 1

    # Total
    ws.cell(r, 1, 9).font = total_font
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(r, 2, "Total Cost").font = total_font
    ws.cell(r, 4, "").font = total_font
    ws.cell(r, 5, "").font = total_font
    for c in range(1, 6):
        ws.cell(r, c).fill = total_fill
        ws.cell(r, c).border = thin
        ws.cell(r, c).alignment = center if c != 2 else left

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[dash_title_row].height = 24
    for rr in range(first_dash_data, first_dash_data + len(dash_rows)):
        ws.row_dimensions[rr].height = 30

    out = Path(__file__).resolve().parents[1] / "NHA_AWS_and_Dashboard_Server_Requirement.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
